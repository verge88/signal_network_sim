#!/usr/bin/env python3
"""Build a self-contained Russian scientific DOCX from the editable manuscript.
The embedded figures correspond to the Mermaid sources under diagrams/.
Graphviz is used for robust offline rendering of the same flowchart topology.
"""
from __future__ import annotations

import re
import subprocess
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "5g_nrf_kie_article_draft.docx"
SOURCE = ROOT / "manuscript.md"
FIGS = ROOT / "figures"
FIGS.mkdir(exist_ok=True)

CAPTIONS = {
    "architecture": "Рис. 1. Разделение пакетного свидетельства и подписанной телеметрии при оценке событий NRF",
    "decision": "Рис. 2. Последовательность проверки внешнего события и отклика KIE",
    "design": "Рис. 3. Схема проспективной серии независимых развертываний Open5GS",
    "results": "Рис. 4. Тревоги KIE v1 и v2 при коротких скрытых воздействиях (по шесть воздействий на период)",
}
DIAGRAM_PATHS = {
    "architecture": "01_architecture",
    "decision": "02_decision",
    "design": "03_design",
    "results": "04_results",
}

def render_diagrams() -> None:
    for key in ("architecture", "decision", "design"):
        stem = DIAGRAM_PATHS[key]
        subprocess.run([
            "dot", "-Tpng", "-Gdpi=170",
            str(ROOT / "diagrams" / (stem + ".dot")),
            "-o", str(FIGS / (key + ".png")),
        ], check=True)
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 10,
        "axes.titlesize": 12, "axes.labelsize": 10,
    })
    periods = ["1 с", "2 с", "4 с"]
    x = list(range(3))
    fig, ax = plt.subplots(figsize=(8.4, 3.3), dpi=170)
    v1 = [5, 4, 4]
    v2 = [5, 4, 3]
    left = [i - .18 for i in x]
    right = [i + .18 for i in x]
    bars1 = ax.bar(left, v1, width=.35, color="#326fa3", label="KIE v1")
    bars2 = ax.bar(right, v2, width=.35, color="#d3904a", label="KIE v2")
    ax.bar_label(bars1, padding=3)
    ax.bar_label(bars2, padding=3)
    ax.set_xticks(x, periods)
    ax.set_ylim(0, 6.4)
    ax.set_yticks(range(7))
    ax.set_ylabel("Тревоги из 6 воздействий")
    ax.grid(axis="y", color="#e5e9ee", linewidth=.7)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, ncol=2, loc="upper right")
    fig.tight_layout()
    fig.savefig(FIGS / "results.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

def set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_border(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "CCD3DD")
        borders.append(el)
    tcPr.append(borders)

def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instruction, separate, text, end):
        run._r.append(item)

def add_rich(p, s: str) -> None:
    """Minimal inline bold parser, preserving the exact authored text."""
    parts = re.split(r"(\*\*.*?\*\*)", s)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)

def paragraph(doc: Document, s: str, *, style: str | None = None):
    p = doc.add_paragraph(style=style)
    add_rich(p, s)
    return p

def add_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        values = [v.strip() for v in line.strip().strip("|").split("|")]
        if values and all(re.fullmatch(r":?-{3,}:?", v) for v in values):
            continue
        rows.append(values)
    if not rows:
        return
    n = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if i == 0:
                set_cell_shading(cell, "E8EEF6")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.first_line_indent = Cm(0)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8.6 if n >= 5 else 9.5)
                    if i == 0:
                        r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_figure(doc: Document, key: str) -> None:
    path = FIGS / (key + ".png")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(16.4))
    p2 = doc.add_paragraph(style="Caption")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(CAPTIONS[key])

def add_mermaid_appendix(doc: Document) -> None:
    h = doc.add_heading("Исходные схемы Mermaid", level=2)
    h.paragraph_format.space_before = Pt(10)
    for key, stem in DIAGRAM_PATHS.items():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run(stem + ".mmd")
        run.bold = True
        source = (ROOT / "diagrams" / (stem + ".mmd")).read_text(encoding="utf-8")
        for line in source.splitlines():
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(7.7)
            run.font.color.rgb = RGBColor(63, 73, 91)

def main() -> None:
    render_diagrams()
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(.8)
    section.footer_distance = Cm(.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(.7)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)
    for hname, size in (("Heading 1", 12.5), ("Heading 2", 11.5)):
        style = doc.styles[hname]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(32, 49, 71)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(9.5)
    cap.font.italic = True
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.text = "ИССЛЕДОВАНИЕ БЕЗОПАСНОСТИ СИГНАЛЬНОГО ОБМЕНА 5G"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(115, 126, 141)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("С. ")
    add_field(footer, "PAGE")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    in_refs = False
    first = True
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(13)
            r = p.add_run(line[2:])
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            i += 1
            continue
        if line.startswith("## "):
            name = line[3:]
            if name == "ЛИТЕРАТУРА":
                doc.add_page_break()
                in_refs = True
            doc.add_heading(name, level=1)
            i += 1
            continue
        if line.startswith("[[FIG:") and line.endswith("]]"):
            add_figure(doc, line[6:-2])
            i += 1
            continue
        if line.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            add_table(doc, tbl)
            continue
        if line.startswith("FORMULA:"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(9)
            r = p.add_run(line.removeprefix("FORMULA:").strip())
            r.italic = True
            r.font.size = Pt(11)
            i += 1
            continue
        if in_refs and re.match(r"^\d+\.\s", line):
            p = paragraph(doc, line)
            p.paragraph_format.first_line_indent = Cm(-.55)
            p.paragraph_format.left_indent = Cm(.65)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.size = Pt(9.2)
            i += 1
            continue
        chunks = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("## ", "# ", "|", "[[FIG:", "FORMULA:")):
            if in_refs and re.match(r"^\d+\.\s", lines[i].strip()):
                break
            chunks.append(lines[i].strip())
            i += 1
        p = paragraph(doc, " ".join(chunks))
        if chunks[0].startswith("**Таблица"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
        if chunks[0].startswith("**Авторы:"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
    add_mermaid_appendix(doc)
    doc.core_properties.title = "Обнаружение маскируемых всплесков сигнальных запросов к NRF в сети 5G"
    doc.core_properties.subject = "Научная статья: Open5GS, NRF, KIE, независимая пакетная валидация"
    doc.core_properties.keywords = "5G; NRF; HTTP/2; KIE; HMAC; Open5GS"
    doc.save(OUT)
    with zipfile.ZipFile(OUT) as z:
        assert z.testzip() is None
        images = [name for name in z.namelist() if name.startswith("word/media/")]
        assert len(images) == 4, images
    check = Document(OUT)
    assert len(check.tables) == 2, len(check.tables)
    assert len(check.paragraphs) > 100, len(check.paragraphs)
    print(f"DOCX_OK path={OUT} bytes={OUT.stat().st_size} paragraphs={len(check.paragraphs)} tables={len(check.tables)} embedded_figures=4")

if __name__ == "__main__":
    main()
